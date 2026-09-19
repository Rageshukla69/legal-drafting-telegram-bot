"""Shared helpers for DOCX/PDF rendering-parity checks.

The helpers deliberately use an *independent* LibreOffice invocation instead of
the application's own converter, so a bug in the application converter (wrong
paper size, wrong frame, stale temp file) cannot make the parity check pass.

Requires: pymupdf (``import fitz``), see ``requirements-dev.txt``.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

# Devanagari block + Devanagari Extended
_DEVANAGARI = re.compile(r"[\u0900-\u097F\uA8E0-\uA8FF]")
_MATRAS = set("\u093E\u093F\u0940\u0941\u0942\u0943\u0944\u0945\u0946\u0947\u0948\u0949\u094A\u094B\u094C\u094D\u0951\u0952\u0953\u0954\u0962\u0963")


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in ("/usr/bin/soffice", "/usr/local/bin/soffice", "/opt/libreoffice/program/soffice"):
        if Path(candidate).exists():
            return candidate
    return None


def reference_docx_to_pdf(docx_path: Path, out_dir: Path) -> Path | None:
    """Convert a DOCX to PDF with a bare-minimum, self-contained soffice call."""
    soffice = find_soffice()
    if not soffice:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    profile = out_dir / "_reference_profile"
    env = dict(os.environ)
    env["HOME"] = str(out_dir)
    cmd = [
        soffice,
        "--headless",
        "--norestore",
        "--nolockcheck",
        "--nodefault",
        "--nofirststartwizard",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        str(out_dir),
        str(docx_path),
    ]
    subprocess.run(cmd, check=True, timeout=300, capture_output=True, env=env)
    produced = out_dir / (docx_path.stem + ".pdf")
    return produced if produced.exists() else None


def docx_paragraph_texts(docx_path: Path) -> list[str]:
    from docx import Document

    document = Document(str(docx_path))
    return [p.text for p in document.paragraphs]


def docx_footer_texts(docx_path: Path) -> list[str]:
    from docx import Document

    document = Document(str(docx_path))
    out: list[str] = []
    for section in document.sections:
        for part in (section.footer, section.header):
            if part is None:
                continue
            out.extend(p.text for p in part.paragraphs)
    return out


def pdf_page_count(pdf_path: Path) -> int:
    import fitz

    with fitz.open(str(pdf_path)) as doc:
        return doc.page_count


def pdf_page_sizes(pdf_path: Path) -> list[tuple[float, float]]:
    import fitz

    with fitz.open(str(pdf_path)) as doc:
        return [(round(p.rect.width, 1), round(p.rect.height, 1)) for p in doc]


def pdf_text(pdf_path: Path) -> str:
    import fitz

    with fitz.open(str(pdf_path)) as doc:
        return "\n".join(page.get_text() for page in doc)


_PAGE_NUMBER_LINE = re.compile(r"^\s*\d+\s*$", re.MULTILINE)


def pdf_body_text(pdf_path: Path, *, strip_page_numbers: bool = True) -> str:
    text = pdf_text(pdf_path)
    if strip_page_numbers:
        text = _PAGE_NUMBER_LINE.sub("", text)
    return text


def squash(text: str) -> str:
    """Whitespace-insensitive normalisation for text comparison.

    Line wrapping differs between the DOCX paragraph model and PDF text
    extraction, so all whitespace is removed before comparing content.
    """
    return re.sub(r"\s+", "", text)


def pdf_first_line_per_page(pdf_path: Path) -> list[str]:
    """First non-empty text line of every page — used for page-break parity."""
    import fitz

    out: list[str] = []
    with fitz.open(str(pdf_path)) as doc:
        for page in doc:
            for raw in page.get_text().splitlines():
                line = raw.strip()
                if line and not _PAGE_NUMBER_LINE.match(line):
                    out.append(line)
                    break
            else:
                out.append("")
    return out


def unicode_integrity(pdf_path: Path) -> dict[str, int]:
    """Report Unicode damage in the PDF text layer."""
    text = pdf_text(pdf_path)
    replacement = text.count("\ufffd")
    pua = sum(1 for c in text if 0xE000 <= ord(c) <= 0xF8FF or 0xF0000 <= ord(c) <= 0xFFFFD)
    control = sum(1 for c in text if ord(c) < 32 and c not in "\n\r\t")
    combining_orphans = 0
    # A matra directly after whitespace/start means its base consonant was lost.
    for match in re.finditer(r"(^|\s)([%s])" % re.escape("".join(sorted(_MATRAS))), text):
        combining_orphans += 1
    devanagari = len(_DEVANAGARI.findall(text))
    return {
        "replacement_chars": replacement,
        "private_use_chars": pua,
        "control_chars": control,
        "orphan_matra_sequences": combining_orphans,
        "devanagari_chars": devanagari,
        "total_chars": len(text),
    }


def contains_replacement_or_pua(text: str) -> bool:
    return any(
        c == "\ufffd" or 0xE000 <= ord(c) <= 0xF8FF or 0xF0000 <= ord(c) <= 0xFFFFD
        for c in text
    )


def ink_difference(pdf_a: Path, pdf_b: Path, *, dpi: int = 110) -> tuple[int, float]:
    """Rasterise two PDFs and return (pages_compared, differing_pixel_ratio).

    Page dimensions must match. Returns (0, 1.0) when they do not.
    """
    import fitz

    with fitz.open(str(pdf_a)) as da, fitz.open(str(pdf_b)) as db:
        if da.page_count != db.page_count:
            return (0, 1.0)
        compared = 0
        total_diff = 0
        total_px = 0
        for i in range(da.page_count):
            pa, pb = da[i], db[i]
            if (round(pa.rect.width), round(pa.rect.height)) != (round(pb.rect.width), round(pb.rect.height)):
                return (0, 1.0)
            matrix = fitz.Matrix(dpi / 72, dpi / 72)
            pix_a = pa.get_pixmap(matrix=matrix, colorspace=fitz.csGRAY)
            pix_b = pb.get_pixmap(matrix=matrix, colorspace=fitz.csGRAY)
            if (pix_a.width, pix_a.height) != (pix_b.width, pix_b.height):
                return (0, 1.0)
            samples_a = pix_a.samples
            samples_b = pix_b.samples
            total_px += len(samples_a)
            compared += 1
            if samples_a == samples_b:
                # Byte-identical rasterisation of this page.
                continue
            total_diff += sum(1 for x, y in zip(samples_a, samples_b) if abs(x - y) > 32)
        return (compared, (total_diff / total_px) if total_px else 1.0)


def devanagari_ratio(text: str) -> float:
    if not text:
        return 0.0
    return len(_DEVANAGARI.findall(text)) / max(1, len(text.strip()))


def is_combining_only_break(text: str) -> bool:
    """True when ``text`` starts with a combining mark (broken grapheme)."""
    stripped = text.lstrip()
    return bool(stripped) and unicodedata.category(stripped[0]) in {"Mn", "Mc"}
