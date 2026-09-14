import os
from pathlib import Path

os.environ.setdefault("GEMINI_API_KEY", "test")

from app.drafting_engine.draft_editor import DraftEditor
from app.drafting_engine.renderers.legal_document_renderer import render_both


def sample_draft():
    return {
        "court_heading": "माननीय न्यायालय सिविल जज",
        "case_heading": "रवि कुमार बनाम राजेश शर्मा",
        "parties": ["वादी: रवि कुमार", "प्रतिवादी: राजेश शर्मा"],
        "title": "रजिस्टर्ड कानूनी नोटिस",
        "opening_averment": "आपको निम्नानुसार सूचित किया जाता है-",
        "pleadings": ["दिनांक 15 फरवरी 2025 को समझौता निष्पादित हुआ।", "प्रतिवादी ने भुगतान वापस नहीं किया।"],
        "prayer": ["प्रतिवादी को 10,00,000 रुपये वापस करने का निर्देश दिया जाए।"],
        "verification": "सत्यापन कथन।",
        "signature_block": ["नोटिसरद्दी- रवि कुमार", "द्वारा अधिवक्ता- वी०डी० शुक्ला"],
    }


def test_natural_language_page_two_edit(monkeypatch):
    editor = DraftEditor()
    editor.client = type("Fake", (), {})()
    edit = {
        "operation": "page_break_before",
        "section": "prayer",
        "index": 0,
        "old_text": "",
        "new_text": "",
        "reason": "User wants prayer on the next page.",
    }
    out, description = editor.apply(sample_draft(), edit)
    assert out["layout"]["page_break_before"] == ["prayer"]
    assert "नए पेज" in description


def test_remove_page_break_is_supported():
    editor = DraftEditor()
    draft = sample_draft()
    draft["layout"] = {"page_break_before": ["prayer"]}
    edit = {
        "operation": "remove_page_break_before",
        "section": "prayer",
        "index": 0,
        "old_text": "",
        "new_text": "",
        "reason": "Remove the page break.",
    }
    out, _ = editor.apply(draft, edit)
    assert out["layout"]["page_break_before"] == []


def test_renderers_honor_prayer_page_break(tmp_path):
    draft = sample_draft()
    draft["layout"] = {"page_break_before": ["prayer"]}
    docx, pdf = render_both(draft, tmp_path, base_name="layout_test", paper="legal")
    assert docx.exists() and pdf.exists()
    from docx import Document
    d = Document(docx)
    prayer_idx = next(i for i,p in enumerate(d.paragraphs) if p.text.strip() == "प्रार्थना")
    assert prayer_idx > 0
    # A page-break paragraph is inserted immediately before the prayer heading.
    assert "w:type=\"page\"" in d.paragraphs[prayer_idx-1]._p.xml or "w:type=\'page\'" in d.paragraphs[prayer_idx-1]._p.xml


def test_screenshot_style_instruction_is_parsed_without_llm():
    editor = DraftEditor()
    editor.client = type("NeverCall", (), {"generate_json": lambda *a, **k: (_ for _ in ()).throw(AssertionError("LLM should not be needed for this clear layout instruction"))})()
    edit = editor.parse(sample_draft(), "इसमें रजिस्ट्री कानूनी नोटिस को पहले पेज पर और जो प्रार्थना है उसको दूसरे पेज पर पहुँचा दो।")
    assert edit["operation"] == "page_break_before"
    assert edit["section"] == "prayer"
    assert edit["old_text"] == "" and edit["new_text"] == ""
