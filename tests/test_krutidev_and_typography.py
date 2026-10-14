from __future__ import annotations

import json
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

from app.drafting_engine.krutidev_converter import decode_krutidev_hybrid, sanitize_draft_unicode
from app.drafting_engine.hindi_corpus_guard import is_clean_hindi_reference, legacy_corruption_score
from app.drafting_engine.hybrid_retriever import HybridCorpusRetriever
from app.drafting_engine.renderers.legal_document_renderer import render_both
from app.drafting_engine.draft_editor import DraftEditor

ROOT = Path(__file__).resolve().parents[1]


def test_krutidev_hybrid_decoder_restores_clean_unicode_hindi():
    corrupted = (
        "न्यायालय श्रhमान सिविल जज जwनियर डिवhजन बिधwना\n"
        "मwल वाद सa[यk सन 2025 ¼दks हजार पPचhस½\n"
        "1& ¼,क½ मqUनh दsवh mमz करhब 71 ¼bकहRतर½ व\"kZ पRनh Lव हरhकृ\"णk\n"
        "mपराsDत वादीगणk निEनलिखित निवsदन करतs हैa&\n"
        "}kरा& वh0डh0 शqDल ,डवाsकेट जिला vौरSयkA"
    )
    cleaned = decode_krutidev_hybrid(corrupted)
    assert "श्रीमान सिविल जज जूनियर डिवीजन बिधूना" in cleaned
    assert "(दो हजार पच्चीस)" in cleaned
    assert "मुन्नी देवी उम्र करीब 71 (इकहत्तर) वर्ष पत्नी स्व हरीकृष्ण" in cleaned
    assert "उपरोक्त वादीगण निम्नलिखित निवेदन करते हैं-" in cleaned
    assert "द्वारा- वी0डी0 शुक्ल एडवोकेट जिला औरैया।" in cleaned
    assert is_clean_hindi_reference(cleaned)


def test_all_306_corpus_documents_are_clean_unicode_and_retrievable():
    index_path = ROOT / "corpus_index.json"
    data = json.loads(index_path.read_text(encoding="utf-8"))
    docs = data["documents"]
    assert len(docs) == 306
    for d in docs:
        text = d.get("text", "")
        if text.strip():
            assert legacy_corruption_score(text) < 0.05, f"Corrupted doc: {d['filename']}"

    retriever = HybridCorpusRetriever(index_path)
    hits = retriever.search("स्थायी निषेधाज्ञा गाटा संख्या बिधूना औरैया", document_type="dava", top_k=3)
    assert len(hits) == 3
    ctx = retriever.format_context(hits)
    assert "ORIGINAL ADVOCATE DRAFT 1" in ctx
    assert "न्यायालय" in ctx


def test_docx_complex_script_bold_and_font_size_tags(tmp_path):
    draft = {
        "court_heading": "न्यायालय श्रीमान सिविल जज जूनियर डिवीजन बिधूना",
        "case_heading": "मूल वाद संख्या 101 सन् 2026",
        "parties": ["वादी: राम कुमार", "प्रतिवादी: श्याम कुमार"],
        "title": "वादपत्र वास्ते स्थायी निषेधाज्ञा",
        "opening_averment": "उपरोक्त वादी निम्नलिखित निवेदन करता है:-",
        "pleadings": ["यह कि विवादित भूमि ग्राम बल्लपुर परगना बिधूना जिला औरैया में स्थित है।"],
        "prayer": ["प्रतिवादी को स्थायी निषेधाज्ञा द्वारा रोका जावे।"],
        "verification": "मैं राम कुमार सत्यापित करता हूँ कि उपरोक्त कथन सत्य है।",
        "signature_block": ["वादी", "द्वारा अधिवक्ता-", "वी०डी० शुक्ला एडवोकेट"],
    }
    docx_path, pdf_path = render_both(draft, tmp_path, base_name="cs_check")
    doc = Document(str(docx_path))
    # First paragraph is court_heading (16pt Bold -> sz=32, szCs=32, b and bCs present)
    court_run = doc.paragraphs[0].runs[0]
    rpr = court_run._element.find(qn("w:rPr"))
    assert rpr.find(qn("w:b")) is not None
    assert rpr.find(qn("w:bCs")) is not None
    assert rpr.find(qn("w:sz")).get(qn("w:val")) == "32"
    assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "32"

    # Body paragraph (14pt Regular -> sz=28, szCs=28, no bCs)
    body_para = next(p for p in doc.paragraphs if "यह कि विवादित" in p.text)
    body_run = body_para.runs[0]
    brpr = body_run._element.find(qn("w:rPr"))
    assert brpr.find(qn("w:bCs")) is None
    assert brpr.find(qn("w:sz")).get(qn("w:val")) == "28"
    assert brpr.find(qn("w:szCs")).get(qn("w:val")) == "28"


def test_draft_editor_replace_all_and_append():
    editor = DraftEditor()
    draft = {
        "court_heading": "न्यायालय श्रीमान सिविल जज जूनियर डिवीजन बिधूना",
        "case_heading": "राम कुमार बनाम श्याम कुमार",
        "parties": ["वादी: राम कुमार", "प्रतिवादी: श्याम कुमार"],
        "title": "वादपत्र वास्ते स्थायी निषेधाज्ञा",
        "opening_averment": "वादी निवेदन करता है:-",
        "pleadings": ["यह कि वादी राम कुमार गाटा संख्या 214 का मालिक है।"],
        "prayer": ["प्रतिवादी को रोका जाए।"],
        "verification": "मैं राम कुमार सत्यापित करता हूँ।",
        "signature_block": ["वादी - राम कुमार"],
    }
    updated, _ = editor.apply(draft, {
        "operation": "replace_all",
        "section": "parties",
        "index": 0,
        "old_text": "राम कुमार",
        "new_text": "राम प्रताप सिंह",
        "reason": "Correct plaintiff name everywhere.",
    })
    assert "राम प्रताप सिंह" in updated["case_heading"]
    assert "राम प्रताप सिंह" in updated["parties"][0]
    assert "राम प्रताप सिंह" in updated["pleadings"][0]
    assert "राम प्रताप सिंह" in updated["verification"]

    appended, _ = editor.apply(updated, {
        "operation": "append",
        "section": "pleadings",
        "index": 0,
        "old_text": "",
        "new_text": "यह कि वाद का कारण दिनांक 10-05-2026 को उत्पन्न हुआ।",
        "reason": "Add cause of action paragraph.",
    })
    assert len(appended["pleadings"]) == 2
