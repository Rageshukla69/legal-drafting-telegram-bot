from pathlib import Path
import tempfile
import zipfile
from docx import Document

from app.drafting_engine.corpus_manager import build_index
from app.drafting_engine.corpus_retriever import retrieve

with tempfile.TemporaryDirectory() as td:
    root = Path(td) / "legal_corpus"
    folder = root / "01_Vaad_Patra_Dava"
    folder.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("न्यायालय श्रीमान सिविल जज, बिधूना")
    doc.add_paragraph("राम कुमार")
    doc.add_paragraph("बनाम")
    doc.add_paragraph("श्याम कुमार")
    doc.add_paragraph("वादपत्र वास्ते स्थायी निषेधाज्ञा")
    doc.add_paragraph("1 (एक) : यह कि वादी कब्जे में है।")
    doc.add_paragraph("प्रार्थना")
    doc.add_paragraph("सत्यापन")
    doc.save(folder / "sample.docx")

    idx = build_index(root)
    assert len(idx["documents"]) == 1
    assert idx["documents"][0]["document_type"] == "dava"
    hits = retrieve(idx, "dava", "कब्जा स्थायी निषेधाज्ञा", 1)
    assert len(hits) == 1

print("PHASE8_SMOKE_OK")
