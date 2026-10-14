```python
"""Deterministic DOCX and PDF renderer for structured legal drafts.

The renderer is deliberately separate from the LLM: the model supplies only
structured content, while this module controls paper size, margins, fonts,
font thickness (both Latin `<w:b/>` and Complex Script `<w:bCs/>`), font sizes
(both `<w:sz/>` and `<w:szCs/>`), alignment, numbering, spacing, and output files.

Canonical pipeline
------------------
The DOCX is the single layout source::

    Structured Draft JSON -> KrutiDev/Unicode Sanitizer -> python-docx (canonical DOCX) -> LibreOffice -> PDF

When running in an environment where LibreOffice is not installed and the
ReportLab fallback is used, `_pdf_inline_font_markup` shapes tokens at the
whitespace-delimited word level so HarfBuzz (`shaping=1`) NEVER switches fonts
mid-token (preventing `.notdef` tofu boxes `□` inside `(एक)` / `(बावन)` and
preventing Greek glyph substitution on commas, hyphens, colons, and parentheses).
"""

from __future__ import annotations

import functools
import logging
import os
import re
import tempfile
from pathlib import Path
from typing import Any
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import LETTER, legal
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageBreak,
    PageTemplate,
    Paragraph,
)

from ..krutidev_converter import sanitize_draft_unicode
from ..legal_number_formatter import format_draft_numbers, hindi_number as _full_hindi_number
from . import docx_to_pdf

log = logging.getLogger(__name__)

HINDI_WORDS = {
    1: "एक", 2: "दो", 3: "तीन", 4: "चार", 5: "पाँच",
    6: "छह", 7: "सात", 8: "आठ", 9: "नौ", 10: "दस",
    11: "ग्यारह", 12: "बारह", 13: "तेरह", 14: "चौदह",
    15: "पंद्रह", 16: "सोलह", 17: "सत्रह", 18: "अठारह",
    19: "उन्नीस", 20: "बीस",
}


def hindi_number(n: int) -> str:
    if n in HINDI_WORDS:
        return HINDI_WORDS[n]
    try:
        return _full_hindi_number(int(n))
    except Exception:
        return str(n)


def numbered(text: str, n: int) -> str:
    text = str(text).strip()
    if not text:
        return ""
    # The model is forbidden from numbering. As a defensive measure, remove
    # one accidental leading numeric/Hindi-letter prefix before rendering.
    text = re.sub(r"^\s*\d+\s*(?:[-–—]?\s*\([^)]*\))?\s*[:.)-]+\s*", "", text)
    return f"{n} ({hindi_number(n)}) : {text}"


def normalize_between_label(value: str) -> str:
    # Dava/Plaint party separator is a deterministic formatting convention.
    return "बनाम"


_BANAM_CLEAN_RE = re.compile(r"^[\s\-–—:|।]*बनाम[\s\-–—:|।]*$")
_LEADING_PARTY_NUM_RE = re.compile(r"^\s*0*(\d+)\s*[.)\-:]\s*")
_ROLE_SUFFIX_RE = re.compile(
    r"(\s*(?:[-–—|।]+\s*)?(?:प्रतिवादीगण|प्रतिवादिनीगण|प्रतिवादिनी|प्रतिवादी|विपक्षीगण|विपक्षी|वादीगण|वादिनी|वादी|प्रार्थीगण|प्रार्थिनी|प्रार्थी)\s*)$"
)


def _sort_numbered_party_lines(lines: list[str]) -> list[str]:
    """Ensure numbered party lines (e.g. '1. शिवमोहन सिंह', '2. श्रीमती रमाकांती') stay in 1, 2, 3 order
    and keep any collective role suffix ('--- प्रतिवादीगण') attached to the final line of the block."""
    if len(lines) <= 1:
        return lines
    numbered_items: list[tuple[int, int, str]] = []
    for idx, line in enumerate(lines):
        m = _LEADING_PARTY_NUM_RE.match(line)
        if m:
            numbered_items.append((int(m.group(1)), idx, line))
    if len(numbered_items) < 2:
        return lines
    nums = [n for n, _, _ in numbered_items]
    if nums == sorted(nums):
        return lines

    # Extract any collective trailing suffix (like '--- प्रतिवादीगण') so it stays on the last line after sorting
    extracted_suffix = ""
    cleaned_lines = list(lines)
    for idx, line in enumerate(cleaned_lines):
        sm = _ROLE_SUFFIX_RE.search(line)
        if sm and idx < len(cleaned_lines) - 1:
            extracted_suffix = sm.group(1)
            cleaned_lines[idx] = line[: sm.start()].rstrip(" ,|-–—")
            break

    sorted_pairs = sorted(
        [
            ((_LEADING_PARTY_NUM_RE.match(l).group(1) if _LEADING_PARTY_NUM_RE.match(l) else "9999"), i, l)
            for i, l in enumerate(cleaned_lines)
        ],
        key=lambda x: (int(x[0]), x[1]),
    )
    result = [l for _, _, l in sorted_pairs]
    if extracted_suffix and result and not _ROLE_SUFFIX_RE.search(result[-1]):
        result[-1] = f"{result[-1]} {extracted_suffix.strip()}"
    return result


def split_party_blocks(parties: list[str]) -> tuple[list[str], list[str], list[str]]:
    """Return plaintiff entries, defendant entries, and any trailing other lines.

    Guarantees:
    - Standalone 'बनाम' entries are stripped and never duplicated.
    - All defendant lines (even when only the last defendant line has '--- प्रतिवादीगण')
      are grouped together in `defendant` in ascending 1, 2, 3 order.
    """
    defendant_labels = (
        "प्रतिवादीगण", "प्रतिवादिनीगण", "प्रतिवादिनी", "प्रतिवादी",
        "विपक्षीगण", "विपक्षी", "संभावित वादीगण", "उत्तरदाता",
        "defendant", "defendants", "respondent", "opposite party",
    )
    plaintiff_labels = (
        "वादीगण", "वादिनी", "वादी", "प्रार्थीगण", "प्रार्थिनी", "प्रार्थी",
        "केवियटकर्तागण", "केवियटकर्ता", "कैवियटकर्तागण", "कैवियटकर्ता",
        "नोटिसदाता", "आवेदक", "आवेदिका",
        "plaintiff", "plaintiffs", "applicant", "petitioner",
    )

    flat_lines: list[str] = []
    for raw in parties or []:
        for line in re.split(r"\r?\n|\\r?\\n", str(raw)):
            s = line.strip()
            if s:
                flat_lines.append(s)

    plaintiff: list[str] = []
    defendant: list[str] = []
    other: list[str] = []

    seen_banam = False
    pending: list[str] = []

    for line in flat_lines:
        if _BANAM_CLEAN_RE.match(line):
            if pending and not plaintiff and not seen_banam:
                plaintiff.extend(pending)
                pending = []
            seen_banam = True
            continue

        low = line.casefold()
        has_def = any(lbl in low for lbl in defendant_labels)
        has_pl = any(lbl in low for lbl in plaintiff_labels)

        if has_def and not has_pl:
            if seen_banam or plaintiff:
                defendant.extend(pending + [line])
            else:
                defendant.extend([line])
            pending = []
        elif has_pl and not has_def:
            if not seen_banam and not defendant:
                plaintiff.extend(pending + [line])
                pending = []
            else:
                plaintiff.append(line)
        else:
            pending.append(line)

    if pending:
        if seen_banam or plaintiff:
            defendant.extend(pending)
        elif defendant:
            defendant.extend(pending)
        else:
            other.extend(pending)

    if not plaintiff and not defendant and len(other) >= 2:
        plaintiff = [other[0]]
        defendant = other[1:]
        other = []

    plaintiff = _sort_numbered_party_lines(plaintiff)
    defendant = _sort_numbered_party_lines(defendant)
    return plaintiff, defendant, other


def _bundled_font(name: str) -> Path | None:
    candidate = Path(__file__).resolve().parents[1] / "assets" / "fonts" / name
    return candidate if candidate.exists() else None


def _font_path() -> str:
    bundled = _bundled_font("NotoSansDevanagari-Regular.ttf")
    if bundled:
        return str(bundled)
    for p in (
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
        "/usr/share/fonts/truetype/lohit-devanagari/Lohit-Devanagari.ttf",
    ):
        if Path(p).exists():
            return p
    raise RuntimeError("Unicode Devanagari font not found.")


def _font_bold_path() -> str:
    bundled = _bundled_font("NotoSansDevanagari-Bold.ttf")
    if bundled:
        return str(bundled)
    for p in (
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Bold.ttf",
        "/usr/share/fonts/truetype/lohit-devanagari/Lohit-Devanagari.ttf",
    ):
        if Path(p).exists():
            return p
    return _font_path()


def _latin_font_path() -> str:
    bundled = _bundled_font("NotoSans-Regular.ttf")
    if bundled:
        return str(bundled)
    bundled_dejavu = _bundled_font("DejaVuSans.ttf")
    if bundled_dejavu:
        return str(bundled_dejavu)
    for p in (
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ):
        if Path(p).exists():
            return p
    raise RuntimeError("Unicode Latin fallback font not found.")


def _latin_bold_font_path() -> str:
    bundled = _bundled_font("DejaVuSans-Bold.ttf")
    if bundled:
        return str(bundled)
    for p in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    ):
        if Path(p).exists():
            return p
    return _latin_font_path()


def _symbol_font_path() -> str:
    bundled = _bundled_font("DejaVuSans.ttf")
    if bundled:
        return str(bundled)
    p = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    if Path(p).exists():
        return p
    return _latin_font_path()


@functools.lru_cache(maxsize=None)
def _font_coverage(path: str) -> frozenset[int]:
    """Return the set of Unicode codepoints a font file can actually render."""
    from fontTools.ttLib import TTFont as _FTFont

    font = _FTFont(path, lazy=True)
    try:
        coverage: set[int] = set()
        for table in font["cmap"].tables:
            coverage.update(table.cmap.keys())
    finally:
        font.close()
    return frozenset(coverage)


@functools.lru_cache(maxsize=None)
def _font_family_name(path: str) -> str:
    """Return the real family name declared inside a font file."""
    from fontTools.ttLib import TTFont as _FTFont

    font = _FTFont(path, lazy=True)
    try:
        for record in font["name"].names:
            if record.nameID == 1 and record.platformID == 3:
                return str(record.toUnicode())
        for record in font["name"].names:
            if record.nameID == 1:
                return str(record.toUnicode())
    finally:
        font.close()
    return Path(path).stem


def _register_pdf_fonts() -> tuple[str, str, str, str]:
    regular = "VakilDeva"
    bold = "VakilDevaBold"
    latin = "VakilLatin"
    latin_bold = "VakilLatinBold"
    symbols = "VakilSymbols"
    registered = pdfmetrics.getRegisteredFontNames()
    if regular not in registered:
        pdfmetrics.registerFont(TTFont(regular, _font_path()))
    if bold not in registered:
        pdfmetrics.registerFont(TTFont(bold, _font_bold_path()))
    if latin not in registered:
        pdfmetrics.registerFont(TTFont(latin, _latin_font_path()))
    if latin_bold not in registered:
        pdfmetrics.registerFont(TTFont(latin_bold, _latin_bold_font_path()))
    if symbols not in registered:
        pdfmetrics.registerFont(TTFont(symbols, _symbol_font_path()))
    return regular, bold, latin, symbols


def _pdf_font_coverage(font_name: str) -> set[int]:
    font = pdfmetrics.getFont(font_name)
    return set(getattr(font.face, "charWidths", {}).keys())


def _is_devanagari(cp: int) -> bool:
    """Return True for Devanagari letters/marks used by Hindi text."""
    return (
        0x0900 <= cp <= 0x097F
        or 0xA8E0 <= cp <= 0xA8FF
        or 0x11B00 <= cp <= 0x11B5F
    )


def _pdf_inline_font_markup(text: str, primary: str, fallback: str, symbols: str) -> str:
    """Create HarfBuzz-safe PDF font runs without ever switching fonts inside a word token.

    CRITICAL FOR REPORTLAB + UHARFBUZZ (`shaping=1`):
    Switching `<font>` tags inside a non-whitespace token (such as `(` in `VakilLatin`
    followed by `एक` in `VakilDeva` in `(एक)`, or `भसौरा` in `VakilDeva` followed by
    `,` in `VakilLatin` in `भसौरा,`) causes ReportLab's HarfBuzz shaper to shape the
    entire non-whitespace cluster with the first font's cmap and draw the second run
    with those glyph IDs in the second font — producing `.notdef` boxes (`□`) for
    parenthesized Hindi words and Greek symbols (`·`, `Έ`, `΄`, `ΗΈ`, `Ό`) for punctuation!

    Because `NotoSansDevanagari` natively covers all ASCII digits `0-9` and all ASCII
    punctuation `( ) , . : ; - / ! ? ' " % + = [ ] { } | ~`, any whitespace-delimited
    token containing Devanagari is kept 100% in `primary` (`VakilDeva`/`VakilDevaBold`).
    """
    primary_cov = _pdf_font_coverage(primary)
    fallback_cov = _pdf_font_coverage(fallback)
    symbol_cov = _pdf_font_coverage(symbols)

    s = str(text or "")
    if not s:
        return ""

    # Separate any direct collision between Latin letters [A-Za-z] and Devanagari,
    # or digits immediately attached to '(' + Devanagari (e.g. '02(रामकान्ती)' -> '02 (रामकान्ती)'),
    # so a single non-whitespace token never contains both scripts.
    s = re.sub(r"(?<=[0-9०-९])\((?=[\u0900-\u097F])", " (", s)
    s = re.sub(r"(?<=[A-Za-z])(?=[\u0900-\u097F])", " ", s)
    s = re.sub(r"(?<=[\u0900-\u097F])(?=[A-Za-z])", " ", s)

    tokens = re.findall(r"\S+|\s+", s)
    out: list[str] = []
    current_font: str | None = None
    buf: list[str] = []

    def flush() -> None:
        nonlocal current_font, buf
        if not buf:
            return
        escaped = _escape_xml("".join(buf))
        if current_font == primary or current_font is None:
            out.append(escaped)
        else:
            out.append(f'<font name="{current_font}">{escaped}</font>')
        buf = []

    def choose_token_font(tok: str) -> str:
        has_deva = any(_is_devanagari(ord(c)) for c in tok)
        if has_deva:
            return primary

        has_ascii_alnum = any(c.isascii() and c.isalnum() for c in tok)
        if has_ascii_alnum:
            return fallback

        needs_fallback = False
        needs_symbol = False
        all_in_primary = True
        for c in tok:
            cp = ord(c)
            cat = unicodedata.category(c)
            if cat in {"Mn", "Mc", "Me", "Cf"}:
                continue
            if cp not in primary_cov:
                all_in_primary = False
                if cp in fallback_cov:
                    needs_fallback = True
                elif cp in symbol_cov:
                    needs_symbol = True

        if needs_symbol and not needs_fallback:
            return symbols
        if needs_fallback or not all_in_primary:
            return fallback
        return primary

    for tok in tokens:
        if tok.isspace():
            buf.append(tok)
            continue

        clean_chars = []
        for ch in tok:
            cp = ord(ch)
            cat = unicodedata.category(ch)
            if cat in {"Mn", "Mc", "Me", "Cf"} or cp in primary_cov or cp in fallback_cov or cp in symbol_cov:
                clean_chars.append(ch)
        if not clean_chars:
            continue
        clean_tok = "".join(clean_chars)

        chosen = choose_token_font(clean_tok)
        if current_font is not None and chosen != current_font:
            # Keep boundary spaces in `primary` (outside `<font name="...">...</font>`)
            if current_font == primary:
                flush()
                current_font = chosen
            else:
                trailing_ws = ""
                if buf and buf[-1].isspace():
                    trailing_ws = buf.pop()
                flush()
                current_font = chosen
                if trailing_ws:
                    if chosen == primary:
                        buf.append(trailing_ws)
                    else:
                        out.append(_escape_xml(trailing_ws))
        else:
            current_font = chosen
        buf.append(clean_tok)

    if current_font is not None and current_font != primary and buf and buf[-1].isspace():
        trailing_ws = buf.pop()
        flush()
        out.append(_escape_xml(trailing_ws))
    else:
        flush()
    return "".join(out)


def _set_run_font(run, name: str, size: float, bold: bool = False):
    """Set font family, size, and bold weight for both Latin and Complex Script (Devanagari)."""
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)

    half_pts = str(int(round(size * 2)))
    sz_el = rpr.find(qn("w:sz"))
    if sz_el is None:
        sz_el = OxmlElement("w:sz")
        rpr.append(sz_el)
    sz_el.set(qn("w:val"), half_pts)

    sz_cs_el = rpr.find(qn("w:szCs"))
    if sz_cs_el is None:
        sz_cs_el = OxmlElement("w:szCs")
        rpr.append(sz_cs_el)
    sz_cs_el.set(qn("w:val"), half_pts)

    b_cs_el = rpr.find(qn("w:bCs"))
    if bold:
        if b_cs_el is None:
            b_cs_el = OxmlElement("w:bCs")
            rpr.append(b_cs_el)
        else:
            b_cs_el.attrib.pop(qn("w:val"), None)
    else:
        if b_cs_el is not None:
            rpr.remove(b_cs_el)


def _docx_font_plan() -> tuple[str, str, str, dict[str, frozenset[int]]]:
    """Resolve the Devanagari / Latin / symbol fonts and their real names."""
    primary_path = _font_path()
    latin_path = _latin_font_path()
    symbol_path = _symbol_font_path()
    coverages = {
        "primary": _font_coverage(primary_path),
        "latin": _font_coverage(latin_path),
        "symbol": _font_coverage(symbol_path),
    }
    return (
        _font_family_name(primary_path),
        _font_family_name(latin_path),
        _font_family_name(symbol_path),
        coverages,
    )


def _add_docx_page_number_footer(document: Document):
    """Put the page number in the DOCX footer as a real `PAGE` field."""
    section = document.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for existing in list(paragraph.runs):
        existing._element.getparent().remove(existing._element)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()
    _set_run_font(run, _font_family_name(_font_path()), 9, bold=False)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, result, end):
        run._element.append(element)
    return paragraph


def _setup_docx(document: Document, paper: str):
    paper = str(paper or "legal").strip().lower()
    if paper not in {"legal", "letter"}:
        paper = "legal"

    section = document.sections[0]
    if paper == "legal":
        section.page_width = Inches(8.5)
        section.page_height = Inches(14)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    try:
        primary_family = _font_family_name(_font_path())
        normal_style = document.styles["Normal"]
        normal_style.font.name = primary_family
        normal_style.font.size = Pt(14)
        rpr = normal_style._element.get_or_add_rPr()
        sz_cs = rpr.find(qn("w:szCs"))
        if sz_cs is None:
            sz_cs = OxmlElement("w:szCs")
            rpr.append(sz_cs)
        sz_cs.set(qn("w:val"), "28")
    except Exception:
        pass

    _add_docx_page_number_footer(document)


def _add_docx_para(
    document: Document,
    text: str,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    size=14,
    bold=False,
    underline=False,
    first_indent=True,
    left_indent=0,
    after=12,
):
    p = document.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing = 1.5
    if first_indent:
        pf.first_line_indent = Inches(0.5)
    if left_indent:
        pf.left_indent = Inches(left_indent)

    text = str(text)
    primary, latin, symbol, coverages = _docx_font_plan()

    runs = []
    current_font = None
    buf = []
    for ch in text:
        cp = ord(ch)
        cat = unicodedata.category(ch)
        if cat in {"Cf", "Mn", "Mc", "Me"}:
            font_name = current_font or primary
        elif ch.isspace() and current_font is not None:
            font_name = current_font
        elif cp in coverages["primary"]:
            font_name = primary
        elif cp in coverages["latin"]:
            font_name = latin
        elif cp in coverages["symbol"]:
            font_name = symbol
        else:
            continue
        if font_name != current_font and buf:
            runs.append((current_font, "".join(buf)))
            buf = []
        current_font = font_name
        buf.append(ch)
    if buf:
        runs.append((current_font, "".join(buf)))
    for font_name, chunk in runs:
        run = p.add_run(chunk)
        _set_run_font(run, font_name, size, bold=bold)
        run.underline = underline
    return p


LAYOUT_SECTIONS = {
    "court_heading", "case_heading", "parties", "title",
    "opening_averment", "pleadings", "prayer", "signature_block", "verification",
}


def _layout_break(draft: dict[str, Any], kind: str, section: str) -> bool:
    layout = draft.get("layout", {}) or {}
    if not isinstance(layout, dict):
        return False
    values = layout.get(kind, []) or []
    if not isinstance(values, list):
        values = [values]
    return section in values


def _docx_page_break_if_needed(doc: Document, draft: dict[str, Any], section: str, *, before: bool = True):
    if _layout_break(draft, "page_break_before" if before else "page_break_after", section):
        doc.add_page_break()


def _ordered_pleading_items(draft: dict[str, Any]) -> list[str]:
    """Return the single approved pleading sequence."""
    return [str(x).strip() for x in (draft.get("pleadings", []) or []) if str(x).strip()]


def _title_text(draft: dict[str, Any]) -> str:
    title = str(draft.get("title", "")).strip()
    if title and title != "वाद पत्र":
        return title
    return "वाद पत्र"


def _prayer_lead_in(draft: dict[str, Any]) -> str:
    title = _title_text(draft)
    if any(k in title for k in ("प्रतिवाद", "जवाब", "लिखित कथन", "Written Statement")):
        return "अतः प्रतिवादी माननीय न्यायालय से प्रार्थना करता है कि:-"
    if any(k in title for k in ("प्रार्थनापत्र", "प्रार्थना पत्र", "आवेदन", "Application")):
        return "अतः श्रीमानजी से प्रार्थना है कि:-"
    return "अतः वादी माननीय न्यायालय से प्रार्थना करता है कि:-"


def _format_prayer_clause(item: str, idx: int) -> str:
    labels = ["(क)", "(ख)", "(ग)", "(घ)", "(ङ)", "(च)", "(छ)", "(ज)", "(झ)", "(ञ)"]
    label = labels[idx] if idx < len(labels) else f"({idx + 1})"
    cleaned = re.sub(r"^\s*(?:\([क-ज्ञ0-9a-zA-Z]+\)|[क-ज्ञअ-ह][.)\-:])\s*", "", str(item).strip())
    return f"{label} {cleaned or item}"


def _signature_lines(draft: dict[str, Any]) -> list[str]:
    lines = [str(x).strip() for x in draft.get("signature_block", []) or [] if str(x).strip()]
    if lines:
        return lines
    plaintiff = ""
    for party in draft.get("parties", []) or []:
        if "वादी" in str(party):
            plaintiff = str(party).strip()
            break
    if plaintiff:
        return [f"वादी\n{plaintiff}", "द्वारा अधिवक्ता-", "वी०डी० शुक्ला एडवोकेट", "बिधूना, औरैया"]
    return ["वादी", "द्वारा अधिवक्ता-", "वी०डी० शुक्ला एडवोकेट", "बिधूना, औरैया"]


def _prepare_draft(draft: dict[str, Any]) -> dict[str, Any]:
    """Sanitize KrutiDev/abbreviations and normalize Hindi years before rendering."""
    return sanitize_draft_unicode(draft or {})


def render_docx(draft: dict[str, Any], output_path: str | Path, paper: str = "legal"):
    draft = _prepare_draft(draft)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _setup_docx(doc, paper)

    court = str(draft.get("court_heading", "")).strip()
    case = str(draft.get("case_heading", "")).strip()
    parties = [str(x).strip() for x in draft.get("parties", []) if str(x).strip()]

    _docx_page_break_if_needed(doc, draft, "court_heading")
    if court:
        _add_docx_para(doc, court, align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, first_indent=False, after=8)
    if case:
        _docx_page_break_if_needed(doc, draft, "case_heading")
        _add_docx_para(doc, case, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, first_indent=False, after=10)

    _docx_page_break_if_needed(doc, draft, "parties")
    plaintiff, defendant, other = split_party_blocks(parties)
    if plaintiff and defendant:
        for party in plaintiff:
            _add_docx_para(doc, party, align=WD_ALIGN_PARAGRAPH.LEFT, size=14, first_indent=False, after=3)
        _add_docx_para(doc, "बनाम", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, first_indent=False, after=3)
        for party in defendant:
            _add_docx_para(doc, party, align=WD_ALIGN_PARAGRAPH.LEFT, size=14, first_indent=False, after=4)
        for party in other:
            _add_docx_para(doc, party, align=WD_ALIGN_PARAGRAPH.LEFT, size=14, first_indent=False, after=4)
    else:
        for party in (plaintiff + defendant + other) or parties:
            _add_docx_para(doc, party, align=WD_ALIGN_PARAGRAPH.LEFT, size=14, first_indent=False, after=4)

    _docx_page_break_if_needed(doc, draft, "title")
    _add_docx_para(doc, _title_text(draft), align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, underline=True, first_indent=False, after=10)

    opening = str(draft.get("opening_averment", "")).strip()
    if opening:
        _docx_page_break_if_needed(doc, draft, "opening_averment")
        _add_docx_para(doc, opening, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, first_indent=False, after=10)

    _docx_page_break_if_needed(doc, draft, "pleadings")
    for idx, paragraph in enumerate(_ordered_pleading_items(draft), 1):
        _add_docx_para(doc, numbered(paragraph, idx), align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, first_indent=True, after=12)

    prayer = [str(x).strip() for x in draft.get("prayer", []) or [] if str(x).strip()]
    if prayer:
        _docx_page_break_if_needed(doc, draft, "prayer")
        _add_docx_para(doc, "प्रार्थना", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, underline=True, first_indent=False, after=8)
        _add_docx_para(doc, _prayer_lead_in(draft), align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, first_indent=False, after=10)
        for idx, item in enumerate(prayer):
            _add_docx_para(doc, _format_prayer_clause(item, idx), align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, first_indent=False, after=10)

    signatures = _signature_lines(draft)
    if signatures:
        _docx_page_break_if_needed(doc, draft, "signature_block")
        for line in signatures:
            _add_docx_para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT, size=14, first_indent=False, after=4)

    verification = str(draft.get("verification", "") or "").strip()
    if verification:
        _docx_page_break_if_needed(doc, draft, "verification")
        _add_docx_para(doc, "सत्यापन", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, underline=True, first_indent=False, after=8)
        for line in verification.splitlines():
            if line.strip():
                _add_docx_para(doc, line.strip(), align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, first_indent=True, after=10)

    doc.save(output_path)
    return output_path


def _escape_xml(text: str) -> str:
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _pdf_story(draft: dict[str, Any], font: str, bold_font: str, latin_font: str, symbol_font: str):
    draft = _prepare_draft(draft)
    latin_bold = "VakilLatinBold" if "VakilLatinBold" in pdfmetrics.getRegisteredFontNames() else latin_font
    styles = getSampleStyleSheet()
    body = ParagraphStyle("VakilBody", parent=styles["BodyText"], fontName=font, fontSize=14, leading=21, alignment=TA_JUSTIFY, shaping=1, firstLineIndent=0.5 * inch, spaceAfter=12)
    center16 = ParagraphStyle("VakilCenter16", parent=body, fontName=bold_font, fontSize=16, leading=22, alignment=TA_CENTER, shaping=1, firstLineIndent=0, spaceAfter=8)
    left = ParagraphStyle("VakilLeft", parent=body, alignment=TA_LEFT, firstLineIndent=0, spaceAfter=3)
    center = ParagraphStyle("VakilCenter", parent=body, fontName=bold_font, alignment=TA_CENTER, firstLineIndent=0, spaceAfter=3, shaping=1)
    right = ParagraphStyle("VakilRight", parent=body, alignment=TA_RIGHT, firstLineIndent=0, spaceAfter=4)
    noindent = ParagraphStyle("VakilNoIndent", parent=body, firstLineIndent=0, spaceAfter=10)
    story = []

    court = str(draft.get("court_heading", "")).strip()
    case = str(draft.get("case_heading", "")).strip()
    parties = [str(x).strip() for x in draft.get("parties", []) if str(x).strip()]
    if _layout_break(draft, "page_break_before", "court_heading"):
        story.append(PageBreak())
    if court:
        story.append(Paragraph(_pdf_inline_font_markup(court, bold_font, latin_bold, symbol_font), center16))
    if case:
        if _layout_break(draft, "page_break_before", "case_heading"):
            story.append(PageBreak())
        story.append(Paragraph(_pdf_inline_font_markup(case, bold_font, latin_bold, symbol_font), ParagraphStyle("Case", parent=center16, fontSize=14, leading=20, spaceAfter=10)))

    if _layout_break(draft, "page_break_before", "parties"):
        story.append(PageBreak())
    plaintiff, defendant, other = split_party_blocks(parties)
    if plaintiff and defendant:
        for party in plaintiff:
            story.append(Paragraph(_pdf_inline_font_markup(party, font, latin_font, symbol_font), left))
        story.append(Paragraph(_pdf_inline_font_markup("बनाम", bold_font, latin_bold, symbol_font), center))
        for party in defendant:
            story.append(Paragraph(_pdf_inline_font_markup(party, font, latin_font, symbol_font), left))
        for party in other:
            story.append(Paragraph(_pdf_inline_font_markup(party, font, latin_font, symbol_font), left))
    else:
        for party in (plaintiff + defendant + other) or parties:
            story.append(Paragraph(_pdf_inline_font_markup(party, font, latin_font, symbol_font), left))

    if _layout_break(draft, "page_break_before", "title"):
        story.append(PageBreak())
    story.append(Paragraph(_pdf_inline_font_markup(_title_text(draft), bold_font, latin_bold, symbol_font), ParagraphStyle("Title", parent=center16, underline=True, spaceBefore=4, spaceAfter=10)))

    opening = str(draft.get("opening_averment", "")).strip()
    if opening:
        if _layout_break(draft, "page_break_before", "opening_averment"):
            story.append(PageBreak())
        story.append(Paragraph(_pdf_inline_font_markup(opening, font, latin_font, symbol_font), noindent))

    if _layout_break(draft, "page_break_before", "pleadings"):
        story.append(PageBreak())
    for idx, paragraph in enumerate(_ordered_pleading_items(draft), 1):
        story.append(Paragraph(_pdf_inline_font_markup(numbered(paragraph, idx), font, latin_font, symbol_font), body))

    prayer = [str(x).strip() for x in draft.get("prayer", []) or [] if str(x).strip()]
    if prayer:
        if _layout_break(draft, "page_break_before", "prayer"):
            story.append(PageBreak())
        story.append(Paragraph("प्रार्थना", ParagraphStyle("PrayerTitle", parent=center16, underline=True, spaceBefore=2, spaceAfter=8)))
        story.append(Paragraph(_pdf_inline_font_markup(_prayer_lead_in(draft), font, latin_font, symbol_font), noindent))
        for idx, item in enumerate(prayer):
            story.append(Paragraph(_pdf_inline_font_markup(_format_prayer_clause(item, idx), font, latin_font, symbol_font), ParagraphStyle(f"Prayer{idx}", parent=body, firstLineIndent=0, spaceAfter=10)))

    if _layout_break(draft, "page_break_before", "signature_block"):
        story.append(PageBreak())
    for line in _signature_lines(draft):
        story.append(Paragraph(_pdf_inline_font_markup(line, font, latin_font, symbol_font), right))

    verification = str(draft.get("verification", "") or "").strip()
    if verification:
        if _layout_break(draft, "page_break_before", "verification"):
            story.append(PageBreak())
        story.append(Paragraph("सत्यापन", ParagraphStyle("VerificationTitle", parent=center16, underline=True, spaceBefore=8, spaceAfter=8)))
        for line in verification.splitlines():
            if line.strip():
                story.append(Paragraph(_pdf_inline_font_markup(line.strip(), font, latin_font, symbol_font), body))
    return story


class NumberedCanvas(Canvas):
    """Adds page numbers without changing the underlying legal content."""

    def __init__(self, *args, **kwargs):
        Canvas.__init__(self, *args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(page_count)
            Canvas.showPage(self)
        Canvas.save(self)

    def draw_page_number(self, page_count: int):
        self.setFont("VakilDeva", 9)
        self.drawCentredString(4.25 * inch, 0.45 * inch, f"{self._pageNumber}")


def render_pdf_reportlab(draft: dict[str, Any], output_path: str | Path, paper: str = "legal"):
    """Fallback PDF layout produced by ReportLab with HarfBuzz token-level font safety."""
    draft = _prepare_draft(draft)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    font, bold_font, latin_font, symbol_font = _register_pdf_fonts()

    paper = str(paper or "legal").strip().lower()
    if paper not in {"legal", "letter"}:
        paper = "legal"

    page_size = legal if paper == "legal" else LETTER
    width, height = page_size
    frame = Frame(
        1.25 * inch, 1 * inch,
        width - 2.5 * inch, height - 2 * inch,
        id="legal_frame",
        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0,
    )
    doc = BaseDocTemplate(
        str(output_path),
        pagesize=page_size,
        leftMargin=1.25 * inch,
        rightMargin=1.25 * inch,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
        title="Legal Draft",
        author="Legal Drafting Bot",
    )
    doc.addPageTemplates([PageTemplate(id="legal", frames=[frame])])
    doc.build(
        _pdf_story(draft, font, bold_font, latin_font, symbol_font),
        canvasmaker=NumberedCanvas,
    )
    return output_path


def pdf_engine() -> str:
    """Return the configured PDF engine: `canonical` (default) or `reportlab`."""
    engine = (os.getenv("LEGAL_PDF_ENGINE", "canonical") or "canonical").strip().lower()
    return engine if engine in {"canonical", "reportlab"} else "canonical"


def _canonical_pdf_required() -> bool:
    return (os.getenv("LEGAL_PDF_REQUIRE_CANONICAL", "") or "").strip().lower() in {
        "1",
        "true",
        "yes",
        "on",
    }


def canonical_pdf_available() -> bool:
    """True when the PDF can be produced from the canonical DOCX."""
    return pdf_engine() == "canonical" and docx_to_pdf.converter_available()


def _write_pdf(
    draft: dict[str, Any],
    pdf_path: str | Path,
    paper: str,
    docx_path: str | Path | None = None,
) -> Path:
    """Write the PDF, preferring a conversion of the canonical DOCX."""
    draft = _prepare_draft(draft)
    pdf_path = Path(pdf_path)
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    if pdf_engine() == "reportlab":
        log.info("LEGAL_PDF_ENGINE=reportlab: using the ReportLab PDF layout.")
        return render_pdf_reportlab(draft, pdf_path, paper=paper)

    if not docx_to_pdf.converter_available():
        message = (
            "LibreOffice (soffice) is not installed, so the PDF cannot be converted from "
            "the canonical DOCX."
        )
        if _canonical_pdf_required():
            raise docx_to_pdf.PdfConverterUnavailable(message)
        log.warning(
            "%s Falling back to the ReportLab layout. Install LibreOffice, or set "
            "SOFFICE_BIN to a LibreOffice binary, to use DOCX->PDF conversion.",
            message,
        )
        return render_pdf_reportlab(draft, pdf_path, paper=paper)

    try:
        if docx_path is not None and Path(docx_path).is_file():
            return docx_to_pdf.convert(docx_path, pdf_path)
        with tempfile.TemporaryDirectory(prefix="legal-canonical-") as tmp:
            staged = Path(tmp) / "canonical.docx"
            render_docx(draft, staged, paper=paper)
            return docx_to_pdf.convert(staged, pdf_path)
    except docx_to_pdf.PdfConversionError:
        log.exception("Canonical DOCX->PDF conversion failed")
        if _canonical_pdf_required():
            raise
        log.warning("Falling back to the ReportLab PDF layout for this document.")
        return render_pdf_reportlab(draft, pdf_path, paper=paper)


def render_pdf(draft: dict[str, Any], output_path: str | Path, paper: str = "legal"):
    """Render the PDF for `draft` from the canonical DOCX."""
    return _write_pdf(draft, output_path, paper)


def render_both(
    draft: dict[str, Any],
    output_dir: str | Path,
    base_name: str = "Dava_Draft",
    paper: str = "legal",
) -> tuple[Path, Path]:
    """Render the DOCX and the PDF, with the PDF converted from that DOCX."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", base_name).strip("_") or "Dava_Draft"
    docx_path = output_dir / f"{safe}.docx"
    pdf_path = output_dir / f"{safe}.pdf"
    render_docx(draft, docx_path, paper=paper)
    _write_pdf(draft, pdf_path, paper=paper, docx_path=docx_path)
    return docx_path, pdf_path
```
